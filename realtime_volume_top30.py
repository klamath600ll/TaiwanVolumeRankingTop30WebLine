#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
台股上市 / 上櫃即時成交量 TOP 30
- 行情來源：TWSE MIS 即時行情服務（盤中資料；非每日收盤行情）
- 股票清單：TWSE / TPEx 公開資料 API
- 單次執行，不輪詢；GitHub Actions 手動啟動
- 成交金額估算 = 最近成交價 × 累積成交股數
  MIS 的累積成交量 v 單位是「張」時，乘以 1,000 換算成股數。
  此數值是以目前價格估算，不是逐筆成交金額加總。
"""

import csv
import io
import json
import os
import sys
import time
from datetime import datetime
from zoneinfo import ZoneInfo

import requests
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

TAIPEI = ZoneInfo("Asia/Taipei")
HTTP_TIMEOUT = 25
BATCH_SIZE = 40
REQUEST_PAUSE = 0.25
TOP_N = 30

TWSE_STOCK_LIST_URL = "https://openapi.twse.com.tw/v1/exchangeReport/STOCK_DAY_ALL"
TPEX_STOCK_LIST_URL = "https://www.tpex.org.tw/openapi/v1/tpex_mainboard_quotes"
MIS_URL = "https://mis.twse.com.tw/stock/api/getStockInfo.jsp"

HEADERS = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
                  "(KHTML, like Gecko) Chrome/124.0 Safari/537.36",
    "Accept": "application/json,text/plain,*/*",
    "Referer": "https://mis.twse.com.tw/stock/index.jsp",
}


def get_json(url, params=None):
    response = requests.get(url, params=params, headers=HEADERS, timeout=HTTP_TIMEOUT)
    response.raise_for_status()
    response.encoding = "utf-8"
    return response.json()


def clean_number(value):
    """將行情字串轉為 float；空值、--、---、除號等回傳 None。"""
    if value is None:
        return None
    value = str(value).strip().replace(",", "")
    if value in ("", "-", "--", "---", "----", "0.00", "0"):
        # 0 可能是有效成交量，但價格為 0 通常代表尚未成交；
        # 這裡仍回傳 0，後續會依欄位判斷。
        if value in ("0", "0.00"):
            return 0.0
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def get_stock_universe():
    """
    取得股票代號 / 名稱清單。
    TWSE 與 TPEx 的這些端點是股票名單來源，不當作盤中即時報價。
    只保留一般數字代號，避免把名稱/代號異常或非股票標的帶入。
    """
    listed, otc = {}, {}

    # TWSE 當日上市證券資料
    twse_rows = get_json(TWSE_STOCK_LIST_URL)
    if not isinstance(twse_rows, list):
        raise RuntimeError("TWSE 股票清單 API 回傳格式非清單。")

    for row in twse_rows:
        code = str(row.get("Code") or row.get("證券代號") or row.get("股票代號") or "").strip()
        name = str(row.get("Name") or row.get("證券名稱") or row.get("股票名稱") or "").strip()
        # 以常見股票代號格式為篩選條件；包含一般股與部分特殊代號，不含權證英數代號
        if code.isdigit() and 4 <= len(code) <= 6 and name:
            listed[code] = name

    # TPEx 上櫃股票資料
    tpex_rows = get_json(TPEX_STOCK_LIST_URL)
    if not isinstance(tpex_rows, list):
        raise RuntimeError("TPEx 股票清單 API 回傳格式非清單。")

    for row in tpex_rows:
        code = str(
            row.get("SecuritiesCompanyCode")
            or row.get("代號")
            or row.get("股票代號")
            or row.get("Code")
            or ""
        ).strip()
        name = str(
            row.get("CompanyName")
            or row.get("名稱")
            or row.get("股票名稱")
            or row.get("Name")
            or ""
        ).strip()
        if code.isdigit() and 4 <= len(code) <= 6 and name:
            # 排除名稱中含「購」「售」的權證類資料，與使用者原本邏輯一致
            if "購" not in name and "售" not in name:
                otc[code] = name

    if not listed:
        raise RuntimeError("未取得上市股票清單；停止執行，避免產生空排行榜。")
    if not otc:
        raise RuntimeError("未取得上櫃股票清單；停止執行，避免產生空排行榜。")

    return listed, otc


def build_exchange_codes(stock_dict, market):
    """MIS ex_ch 格式：tse_2330.tw 或 otc_6488.tw"""
    prefix = "tse" if market == "上市" else "otc"
    return [f"{prefix}_{code}.tw" for code in stock_dict]


def fetch_mis_quotes(ex_ch_codes):
    """
    將股票代號分批送至 MIS。
    每次批次帶多個 ex_ch，避免每檔股票各自發送請求。
    MIS 不是交易所正式授權的即時行情再散布 API；使用可能受服務限制。
    """
    quotes = []
    for start in range(0, len(ex_ch_codes), BATCH_SIZE):
        batch = ex_ch_codes[start:start + BATCH_SIZE]
        params = {
            "ex_ch": "|".join(batch),
            "json": "1",
            "delay": "0",
        }
        try:
            payload = get_json(MIS_URL, params=params)
            msg_array = payload.get("msgArray", []) if isinstance(payload, dict) else []
            if isinstance(msg_array, list):
                quotes.extend(msg_array)
        except requests.RequestException as exc:
            print(f"⚠️ MIS 行情批次失敗（第 {start + 1} 檔起）：{exc}", file=sys.stderr)
        time.sleep(REQUEST_PAUSE)
    return quotes


def quote_to_record(item, stock_dict, market):
    code = str(item.get("c", "")).strip()
    name = str(item.get("n", "")).strip() or stock_dict.get(code, "")
    if not code or not name:
        return None

    # z 最近成交價；若尚未成交，z 可能是 "-"，可用 y 昨收價顯示但不列入成交量排名。
    price = clean_number(item.get("z"))
    volume_lots = clean_number(item.get("v"))
    quote_time = str(item.get("t", "")).strip()
    quote_date = str(item.get("d", "")).strip()

    if volume_lots is None:
        return None
    if volume_lots <= 0:
        return None
    if price is None or price <= 0:
        return None

    # v 為累積成交量（張），估算金額 = 最新成交價 × 累積成交股數
    amount_est = price * volume_lots * 1000

    return {
        "market": market,
        "code": code,
        "name": name,
        "price": price,
        "volume_lots": int(volume_lots),
        "amount_est": amount_est,
        "quote_date": quote_date,
        "quote_time": quote_time,
    }


def fetch_market(market, stock_dict):
    expected = set(stock_dict.keys())
    codes = build_exchange_codes(stock_dict, market)
    raw_quotes = fetch_mis_quotes(codes)

    records = []
    seen = set()
    for item in raw_quotes:
        record = quote_to_record(item, stock_dict, market)
        if not record or record["code"] not in expected or record["code"] in seen:
            continue
        seen.add(record["code"])
        records.append(record)

    records.sort(key=lambda row: row["volume_lots"], reverse=True)
    return records[:TOP_N], len(records), len(stock_dict)


def format_amount(amount):
    if amount >= 1e8:
        return f"{amount / 1e8:,.2f} 億"
    if amount >= 1e4:
        return f"{amount / 1e4:,.2f} 萬"
    return f"{amount:,.0f}"


def write_excel(listed_rows, otc_rows, data_time, output_path):
    wb = Workbook()
    # Create both tabs
    ws_listed = wb.active
    ws_listed.title = "上市成交量TOP30"
    ws_otc = wb.create_sheet("上櫃成交量TOP30")

    headers = ["股票代號", "股票名稱", "股票即時報價", "成交量（張）", "成交金額估算（元）"]
    for ws, rows, title in (
        (ws_listed, listed_rows, "上市"),
        (ws_otc, otc_rows, "上櫃"),
    ):
        ws.append([f"資料時間：{data_time}（台灣時間）｜市場：{title}｜行情來源：TWSE MIS"])
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
        ws["A1"].font = Font(bold=True, color="FFFFFF", size=12)
        ws["A1"].fill = PatternFill("solid", fgColor="1F4E78")
        ws["A1"].alignment = Alignment(horizontal="left")
        ws.append(headers)

        for cell in ws[2]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="4472C4")
            cell.alignment = Alignment(horizontal="center")

        for row in rows:
            ws.append([
                row["code"],
                row["name"],
                row["price"],
                row["volume_lots"],
                row["amount_est"],
            ])

        for r in range(3, ws.max_row + 1):
            ws.cell(r, 3).number_format = '#,##0.00'
            ws.cell(r, 4).number_format = '#,##0'
            ws.cell(r, 5).number_format = '#,##0'
        widths = [14, 20, 18, 18, 25]
        for idx, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(idx)].width = width
        ws.freeze_panes = "A3"
        ws.auto_filter.ref = f"A2:E{ws.max_row}"

    wb.save(output_path)


def write_csv(rows, market, data_time, output_path):
    with open(output_path, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        writer.writerow([f"資料時間：{data_time}（台灣時間）｜市場：{market}｜行情來源：TWSE MIS"])
        writer.writerow(["股票代號", "股票名稱", "股票即時報價", "成交量（張）", "成交金額估算（元）"])
        for row in rows:
            writer.writerow([
                row["code"], row["name"], row["price"], row["volume_lots"], round(row["amount_est"])
            ])



def write_web_page(listed_rows, otc_rows, data_time, site_dir="site"):
    """產生 GitHub Pages 靜態網頁；不將任何 Token 寫入頁面。"""
    import html
    from pathlib import Path

    site_path = Path(site_dir)
    site_path.mkdir(parents=True, exist_ok=True)

    def render_rows(rows):
        if not rows:
            return '<tr><td colspan="5" class="empty">目前沒有可用的成交行情</td></tr>'
        rendered = []
        for rank, row in enumerate(rows, start=1):
            rendered.append(
                "<tr>"
                f'<td class="rank">{rank}</td>'
                f'<td>{html.escape(str(row["code"]))}</td>'
                f'<td>{html.escape(str(row["name"]))}</td>'
                f'<td class="num">{row["price"]:,.2f}</td>'
                f'<td class="num">{row["volume_lots"]:,}</td>'
                f'<td class="num">{row["amount_est"]:,.0f}</td>'
                "</tr>"
            )
        return "\n".join(rendered)

    page = f"""<!doctype html>
<html lang="zh-Hant">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<meta name="robots" content="noindex,nofollow">
<title>台股即時成交量 TOP 30</title>
<style>
:root {{ color-scheme: light; --ink:#172033; --muted:#64748b; --line:#dbe3ee; --blue:#174a83; --bg:#f3f6fa; }}
* {{ box-sizing:border-box; }}
body {{ margin:0; background:var(--bg); color:var(--ink); font-family:system-ui,-apple-system,"Noto Sans TC",sans-serif; }}
header {{ background:linear-gradient(120deg,#102d50,#1e6091); color:white; padding:24px 16px; }}
header .wrap, main {{ max-width:1180px; margin:auto; }}
h1 {{ font-size:clamp(1.35rem,4vw,2rem); margin:0 0 8px; }}
.meta {{ color:#dbeafe; font-size:.92rem; line-height:1.7; }}
main {{ padding:18px 14px 36px; }}
.note {{ background:#fff8e7; border:1px solid #f2d58a; border-radius:10px; padding:12px 14px; color:#6b4e10; font-size:.9rem; line-height:1.65; margin-bottom:18px; }}
.tabs {{ display:flex; gap:8px; flex-wrap:wrap; margin:0 0 14px; }}
button {{ border:1px solid var(--line); border-radius:999px; padding:10px 18px; background:white; color:var(--ink); font-weight:700; cursor:pointer; }}
button.active {{ background:var(--blue); color:white; border-color:var(--blue); }}
.panel {{ display:none; }}
.panel.active {{ display:block; }}
h2 {{ font-size:1.2rem; margin:8px 0 12px; }}
.table-wrap {{ overflow-x:auto; background:white; border:1px solid var(--line); border-radius:12px; box-shadow:0 4px 16px #0f172a0a; }}
table {{ border-collapse:collapse; width:100%; min-width:760px; font-size:.94rem; }}
th {{ background:#eaf1f8; color:#243b53; text-align:left; white-space:nowrap; }}
th,td {{ padding:11px 12px; border-bottom:1px solid var(--line); }}
tbody tr:nth-child(even) {{ background:#f8fafc; }}
tbody tr:last-child td {{ border-bottom:0; }}
.num {{ text-align:right; font-variant-numeric:tabular-nums; white-space:nowrap; }}
.rank {{ color:var(--muted); width:54px; }}
.empty {{ text-align:center; color:var(--muted); padding:28px; }}
footer {{ color:var(--muted); font-size:.82rem; margin-top:16px; line-height:1.6; }}
@media(max-width:600px) {{ header {{ padding:20px 14px; }} main {{ padding:14px 10px 28px; }} }}
</style>
</head>
<body>
<header>
  <div class="wrap">
    <h1>台股即時成交量 TOP 30</h1>
    <div class="meta">資料擷取時間：{html.escape(data_time)}（台灣時間）<br>
    行情來源：TWSE MIS｜上市、上櫃分開排名</div>
  </div>
</header>
<main>
  <div class="note">
    成交金額為估算值：最近成交價 × 當日累積成交量（張）× 1,000。
    此估算並非逐筆成交金額加總。各股票最近成交時間可能不同；非交易時段可能顯示最後可取得行情。
  </div>
  <div class="tabs" role="tablist">
    <button class="active" onclick="showPanel('listed',this)">上市 TOP 30</button>
    <button onclick="showPanel('otc',this)">上櫃 TOP 30</button>
  </div>
  <section id="listed" class="panel active">
    <h2>上市成交量排名</h2>
    <div class="table-wrap"><table>
      <thead><tr><th>排名</th><th>股票代號</th><th>股票名稱</th><th class="num">即時報價（元）</th><th class="num">成交量（張）</th><th class="num">成交金額估算（元）</th></tr></thead>
      <tbody>{render_rows(listed_rows)}</tbody>
    </table></div>
  </section>
  <section id="otc" class="panel">
    <h2>上櫃成交量排名</h2>
    <div class="table-wrap"><table>
      <thead><tr><th>排名</th><th>股票代號</th><th>股票名稱</th><th class="num">即時報價（元）</th><th class="num">成交量（張）</th><th class="num">成交金額估算（元）</th></tr></thead>
      <tbody>{render_rows(otc_rows)}</tbody>
    </table></div>
  </section>
  <footer>
    本頁由 GitHub Actions 依手動執行結果更新。頁面公開可讀，請勿放入個人資料或 API Token。
  </footer>
</main>
<script>
function showPanel(id,button) {{
  document.querySelectorAll('.panel').forEach(p=>p.classList.remove('active'));
  document.querySelectorAll('.tabs button').forEach(b=>b.classList.remove('active'));
  document.getElementById(id).classList.add('active');
  button.classList.add('active');
}}
</script>
</body>
</html>"""
    (site_path / "index.html").write_text(page, encoding="utf-8")
    print(f"GitHub Pages 網頁已產生：{site_path / 'index.html'}")


def send_line_link(site_url):
    """使用 LINE Messaging API 推播網頁連結。Secrets 缺漏或 API 失敗時明確報錯。"""
    import requests

    token = os.environ.get("LINE_CHANNEL_ACCESS_TOKEN", "").strip()
    user_id = os.environ.get("LINE_USER_ID", "").strip()
    if not token or not user_id:
        raise RuntimeError(
            "缺少 LINE_CHANNEL_ACCESS_TOKEN 或 LINE_USER_ID GitHub Secret。"
        )

    endpoint = "https://api.line.me/v2/bot/message/push"
    headers = {
        "Authorization": f"Bearer {token}",
        "Content-Type": "application/json",
    }
    payload = {
        "to": user_id,
        "messages": [{
            "type": "text",
            "text": f"台股即時成交量 TOP 30 已更新\\n{site_url}"
        }]
    }
    response = requests.post(endpoint, headers=headers, json=payload, timeout=HTTP_TIMEOUT)
    if not response.ok:
        # 不輸出 Token，僅列出 HTTP 狀態與 API 回應內容
        raise RuntimeError(f"LINE 推播失敗：HTTP {response.status_code} {response.text}")
    print("LINE 網頁連結通知已送出。")


def main():
    now = datetime.now(TAIPEI)
    run_time = now.strftime("%Y-%m-%d %H:%M:%S")
    print(f"開始查詢台股即時成交量排行榜｜程式啟動時間：{run_time}（台灣時間）")
    print("提醒：資料時間以行情來源回傳的交易日期/最近成交時間為準；程式啟動時間不等於每檔股票成交時間。")

    listed_dict, otc_dict = get_stock_universe()
    print(f"股票清單：上市 {len(listed_dict)} 檔；上櫃 {len(otc_dict)} 檔")

    listed_top, listed_count, listed_total = fetch_market("上市", listed_dict)
    otc_top, otc_count, otc_total = fetch_market("上櫃", otc_dict)

    # 各行情回傳時間不一定完全一致。使用執行時刻作為報表擷取時間，
    # 同時在報表中明確標示是擷取時間，而不是交易所統一時間戳。
    data_time = run_time

    if not listed_top and not otc_top:
        raise RuntimeError(
            "MIS 未回傳可用的成交行情。可能是非交易時段、服務限制或資料來源異常；"
            "不輸出空白排行榜。"
        )

    out_xlsx = os.environ.get("OUTPUT_XLSX", "taiwan_stock_volume_top30.xlsx")
    write_excel(listed_top, otc_top, data_time, out_xlsx)
    write_csv(listed_top, "上市", data_time, "listed_volume_top30.csv")
    write_csv(otc_top, "上櫃", data_time, "otc_volume_top30.csv")
    write_web_page(listed_top, otc_top, data_time, site_dir="site")

    print(f"上市：掃描股票 {listed_total} 檔，取得有成交行情 {listed_count} 檔，輸出 {len(listed_top)} 檔")
    print(f"上櫃：掃描股票 {otc_total} 檔，取得有成交行情 {otc_count} 檔，輸出 {len(otc_top)} 檔")
    print(f"Excel 已輸出：{out_xlsx}")
    print("CSV 已輸出：listed_volume_top30.csv、otc_volume_top30.csv")
    print("成交金額估算公式：最新成交價 × 累積成交量（張）× 1,000。")


if __name__ == "__main__":
    main()


# Workflow 會在 Pages 部署完成後，另行呼叫 send_line_link(site_url)。
